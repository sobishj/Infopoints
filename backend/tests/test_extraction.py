"""Chunk location metadata: page / slide / heading+lines, OCR of scanned pages, and no chunk crossing a page."""
import io
import os

import pymupdf as fitz
import pytest

from tests.conftest import make_pdf
from worker.chunker import chunk_segments
from worker.extract.office import convert_to_pdf
from worker.extract.pdf import extract_pdf
from worker.extract.text import extract_markdown, extract_text

LONG = " ".join(f"Sentence {i} explains the approval workflow for purchase orders in detail." for i in range(200))


def test_pdf_pages_and_chunks_keep_page_numbers(tmp_path):
    path = str(tmp_path / "guide.pdf")
    make_pdf(path, ["Introduction to the procurement module.", LONG[:3000], "Approve: Procurement > Pending Orders > Approve"])
    ex = extract_pdf(path)
    assert ex.page_count == 3
    assert [s.page for s in ex.segments] == [1, 2, 3]
    chunks = chunk_segments(ex.segments, "guide.pdf")
    assert {c.segment.page for c in chunks} == {1, 2, 3}
    # Every chunk's text comes from exactly one page.
    page_text = {s.page: s.text for s in ex.segments}
    for c in chunks:
        first_words = c.text.split()[:5]
        assert " ".join(first_words) in " ".join(page_text[c.segment.page].split())
    approve = [c for c in chunks if "Pending Orders" in c.text]
    assert approve and approve[0].header == "[Source: guide.pdf | Page 3]"


def test_long_page_is_split_with_overlap_but_stays_on_its_page(tmp_path):
    path = str(tmp_path / "long.pdf")
    make_pdf(path, ["short", LONG[:3500]])
    chunks = chunk_segments(extract_pdf(path).segments, "long.pdf")
    page2 = [c for c in chunks if c.segment.page == 2]
    assert len(page2) >= 2
    assert all(c.token_count <= 560 for c in page2)


def test_scanned_page_is_ocrd(tmp_path):
    # Render a text page to an image, then build a PDF whose page 2 contains only that image.
    src = fitz.open()
    p = src.new_page()
    p.insert_text((72, 120), "Invoice Approval Screen", fontsize=28)
    png = p.get_pixmap(dpi=200).tobytes("png")
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "First page has normal text that is long enough to skip OCR entirely.")
    doc.new_page().insert_image(fitz.Rect(0, 0, 595, 842), stream=png)
    path = str(tmp_path / "scan.pdf")
    doc.save(path)
    ex = extract_pdf(path)
    page2 = [s for s in ex.segments if s.page == 2]
    assert page2 and "approval" in page2[0].text.lower()


def test_markdown_headings_and_line_ranges(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("# Setup\nInstall the client.\n\n## Approvals\nGo to Procurement.\nClick Approve.\n", encoding="utf-8")
    ex = extract_markdown(str(path))
    assert [(s.heading, s.line_start, s.line_end) for s in ex.segments] == [("Setup", 1, 3), ("Approvals", 4, 6)]
    chunks = chunk_segments(ex.segments, "notes.md")
    assert chunks[1].header == "[Source: notes.md | Section “Approvals”, lines 4–6]"


def test_text_line_windows(tmp_path):
    path = tmp_path / "log.txt"
    path.write_text("\n".join(f"line {i}" for i in range(1, 131)), encoding="utf-8")
    ex = extract_text(str(path))
    assert [(s.line_start, s.line_end) for s in ex.segments] == [(1, 60), (61, 120), (121, 130)]


@pytest.mark.slow
def test_docx_converted_pages_are_real_page_numbers(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.enum.text import WD_BREAK
    d = docx.Document()
    d.add_paragraph("Chapter one: logging in to the portal.")
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("Chapter two: the vendor onboarding checklist.")
    src = tmp_path / "manual.docx"
    d.save(str(src))
    pdf = convert_to_pdf(str(src), str(tmp_path / "out"))
    ex = extract_pdf(pdf)
    by_page = {s.page: s.text for s in ex.segments}
    assert "logging in" in by_page[1] and "vendor onboarding" in by_page[2]


@pytest.mark.slow
def test_pptx_slides(tmp_path):
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation()
    for title in ("Sprint 12 demo", "Approval flow walkthrough"):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
    src = tmp_path / "demo.pptx"
    prs.save(str(src))
    pdf = convert_to_pdf(str(src), str(tmp_path / "out"))
    ex = extract_pdf(pdf, loc_type="slide")
    chunks = chunk_segments(ex.segments, "demo.pptx")
    hit = [c for c in chunks if "Approval flow" in c.text][0]
    # LibreOffice exports slide titles as PDF bookmarks, so each slide carries its title as the section.
    assert hit.segment.page == 2 and hit.header == "[Source: demo.pptx | Slide 2 | Approval flow walkthrough]"
